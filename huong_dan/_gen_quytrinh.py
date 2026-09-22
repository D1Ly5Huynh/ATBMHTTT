#!/usr/bin/env python3
"""Xuất Word: full quy trình tái hiện nhánh A + nhánh B."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

OUT = Path(__file__).resolve().parent / "Quy_trinh_tai_hien_day_du.docx"
NAVY = RGBColor(0x1F, 0x4E, 0x79)
RED = RGBColor(0x9B, 0x2C, 0x2C)
GREEN = RGBColor(0x1F, 0x6B, 0x4A)


def rfonts(run, name="Times New Roman"):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)


def set_run(run, text, size=13, bold=False, italic=False, color=None, font="Times New Roman"):
    run.text = text
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    rfonts(run, font)


def shade_p(p, fill="F2F2F2"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def shade_cell(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def para(doc, text="", size=13, bold=False, italic=False, color=None, align="justify", space_after=8, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    if text:
        r = p.add_run()
        set_run(r, text, size=size, bold=bold, italic=italic, color=color)
    return p


def add_mixed(p, parts):
    """parts: list of str or (text, kwargs)."""
    for item in parts:
        if isinstance(item, str):
            r = p.add_run()
            set_run(r, item)
        else:
            text, kw = item
            r = p.add_run()
            set_run(r, text, **kw)


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        rfonts(run, "Times New Roman")
        run.font.color.rgb = NAVY
        run.font.size = Pt(16 if level == 1 else 14 if level == 2 else 13)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8)
    return p


def code(doc, text):
    """Khối lệnh, từng dòng Courier New 10pt, nền xám."""
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(4 if i == 0 else 0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm(0.3)
        shade_p(p, "F4F4F4")
        r = p.add_run()
        set_run(r, line if line else " ", size=10, font="Courier New")
    # spacer
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(8)
    sp.paragraph_format.space_before = Pt(2)


def kv(doc, text):
    p = para(doc, "", space_after=8)
    shade_p(p, "E2EFDA")
    add_mixed(p, [
        ("Kỳ vọng: ", {"bold": True, "color": GREEN, "size": 12}),
        (text, {"size": 12}),
    ])


def cam(doc, text):
    p = para(doc, "", space_after=8)
    shade_p(p, "FCE4D6")
    add_mixed(p, [
        ("Cấm: ", {"bold": True, "color": RED, "size": 12}),
        (text, {"size": 12}),
    ])


def tbl(doc, header, rows, col_w=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for j, htxt in enumerate(header):
        c = t.rows[0].cells[j]
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        set_run(r, htxt, size=11, bold=True, color=RGBColor(255, 255, 255))
        shade_cell(c, "1F4E79")
        set_cell_border(c)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.rows[i + 1].cells[j]
            c.text = ""
            p = c.paragraphs[0]
            r = p.add_run()
            set_run(r, str(val), size=11, bold=(j == 0))
            shade_cell(c, "F2F2F2" if i % 2 else "FFFFFF")
            set_cell_border(c)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    return t


def footer_page(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    set_run(r, "ATBMHTTT  ·  Quy trình tái hiện  ·  Trang ", size=10, color=NAVY)
    # PAGE field
    r2 = p.add_run()
    rfonts(r2, "Times New Roman")
    r2.font.size = Pt(10)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r2._r.append(fld1)
    r2._r.append(instr)
    r2._r.append(fld2)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    footer_page(sec)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    rPr = style.element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Times New Roman")

    # ----- COVER -----
    para(doc, "HỌC PHẦN AN TOÀN BẢO MẬT HỆ THỐNG THÔNG TIN (ATBMHTTT)", size=14, bold=True, color=NAVY, align="center", space_after=4)
    para(doc, "QUY TRÌNH TÁI HIỆN ĐẦY ĐỦ", size=22, bold=True, color=NAVY, align="center", space_before=24, space_after=6)
    para(doc, "Nhánh A (sát paper Pasini) và nhánh B (lab JSDOM / CRS)", size=14, italic=True, align="center", space_after=12)
    para(doc, "Repo: /home/kali/Desktop/ATBMHTTT", align="center", size=12, space_after=4)
    para(doc, "Artifact pin a299bb6  ·  Paper: Pasini et al., JSS 2026, arXiv:2502.19095v2", align="center", size=12, space_after=16)
    para(
        doc,
        "Tài liệu này liệt kê đủ lệnh, thư mục làm việc, kỳ vọng số, và điều cấm để dựng lại đồ án từ máy trống. "
        "Mọi lệnh Python dùng .venv/bin/python. Hai nhánh không trộn số, không trộn encode, không trộn Oracle.",
        space_after=12,
    )

    h(doc, "1. Hai nhánh — đọc trước khi gõ lệnh", 1)
    para(
        doc,
        "Đồ án có hai pipeline độc lập, cùng 27 action Table 2 của Chen/Pasini. Kết quả ER ~99% chỉ xuất hiện trên nhánh A. "
        "Nhánh B khép token-id thì ER rơi còn khoảng 1%. Viết báo cáo phải tách bảng.",
    )
    tbl(
        doc,
        ["", "Nhánh B — luận văn (lab/)", "Nhánh A — sát paper (artifact/)"],
        [
            ["Dataset", "Mereani, lọc JSDOM execute", "CSV data/10 có sẵn, không chạy prepare_dataset.py"],
            ["Encode", "vocab.json cố định", "LabelEncoder.fit_transform từng payload"],
            ["Oracle", "JSDOM hook JS + Playwright hold-out", "FastAPI :5555, BeautifulSoup + zss vs \"abc\""],
            ["Agent", "PPO (P4 thêm Dueling DQN)", "PPO SB3, 250.000 bước"],
            ["ER PPO điển hình", "~0,75% → 0,25%", "~99%"],
            ["CRS / DOMPurify", "Có, TASR = 0", "Không dùng"],
            ["File số", "lab/data/P1…P5, PPO_VS_DDQN.md", "lab/data/PASINI_FAITHFUL_RESULTS.md"],
        ],
    )
    cam(doc, "Gọi HTTP 200 là TASR; train RL trên CRS; LLM sinh payload; đánh WAF cloud; import mutators.py vào train nhánh B; chạy prepare_dataset.py (mất split paper).")

    h(doc, "2. Máy cần gì", 1)
    tbl(
        doc,
        ["Thành phần", "Tối thiểu", "Ghi chú"],
        [
            ["OS", "Linux (Kali / Ubuntu 22.04+)", "Đã kiểm trên Kali"],
            ["RAM", "8 GB", "Docker + PPO"],
            ["Disk", "~8 GB", "node_modules, torch, image CRS"],
            ["CPU", "4 luồng", "Không cần GPU, PPO device=cpu"],
            ["Python", "3.11–3.13", "venv ở gốc repo"],
            ["Node.js", "≥ 20", "JSDOM oracle"],
            ["Docker", "24+ + compose", "app :13000, CRS :18080/:18081"],
            ["Git", "clone artifact", "commit a299bb6"],
        ],
    )
    para(doc, "Cổng chỉ bind 127.0.0.1: app 13000, CRS PL1 18080, CRS PL2 18081, Oracle Pasini 5555 (chỉ khi chạy nhánh A).")

    h(doc, "3. Cài đặt chung (làm một lần)", 1)
    h(doc, "3.1. Kiểm tra cây repo", 2)
    code(doc, """cd /home/kali/Desktop/ATBMHTTT
ls lab/docker-compose.yml lab/src/run_p3.py lab/data/Payloads.csv
ls artifact/Adversarial_RL_XSS/data/10/detectors/train.csv""")
    para(doc, "Nếu chưa có artifact:")
    code(doc, """mkdir -p artifact
git clone https://github.com/GianlucaMaragliano/Adversarial_RL_XSS.git artifact/Adversarial_RL_XSS
git -C artifact/Adversarial_RL_XSS checkout a299bb6""")
    para(doc, "Payloads.csv encoding latin-1. Không cần FigShare nếu data/10 đã có 7 CSV.")

    h(doc, "3.2. Python venv", 2)
    code(doc, """cd /home/kali/Desktop/ATBMHTTT
python3 -m venv .venv
source .venv/bin/activate
pip install -U pip
pip install -r lab/requirements.txt
pip install scipy scikit-learn nltk validators zss fastapi uvicorn \\
  jinja2 python-dotenv pydantic-settings requests markupsafe
.venv/bin/python -c "import torch, pandas, gymnasium, stable_baselines3; print('py_ok', torch.__version__)" """)
    para(doc, "lab/requirements.txt: beautifulsoup4, pandas, lxml, torch>=2.2, stable-baselines3>=2.3, gymnasium>=1.0. Torch CPU đủ.")

    h(doc, "3.3. Node — JSDOM + Playwright", 2)
    code(doc, """cd /home/kali/Desktop/ATBMHTTT/lab
npm install
npx playwright install chromium
npx playwright install-deps chromium""")
    para(doc, "App Docker tự npm install express + DOMPurify. Không trộn với lab/node_modules.")

    h(doc, "4. NHÁNH B — lab/ (JSDOM, vocab cố định)", 1)
    para(doc, "Thứ tự bắt buộc: Oracle → lọc Mereani → split → 27 action → Docker/TASR → LSTM+PPO → Oracle-in-loop. Sai kỳ vọng thì dừng, đừng train PPO.")

    h(doc, "4.1. B1 — Smoke Oracle và unit test", 2)
    code(doc, """cd /home/kali/Desktop/ATBMHTTT/lab
node oracle/jsdom_oracle.mjs --payload '<script>alert(1)</script>'
node oracle/jsdom_oracle.mjs --payload '<img src=x onerror=alert(1)>'
node oracle/jsdom_oracle.mjs --payload '<a href="javascript:alert(1)">x</a>'
node oracle/jsdom_oracle.mjs --payload '<div>hello</div>'
../.venv/bin/python -m unittest tests.test_oracle tests.test_actions -v""")
    kv(doc, "Ba payload đầu executed = true. <div>hello</div> executed = false (paper DOM-oracle vẫn gọi độc — đó là W1). Unittest xanh. A18 không no-op; A21 không thành (1)(1).")

    h(doc, "4.2. B2 — P1 lọc Mereani (gate RR = 0)", 2)
    para(doc, "Bỏ qua nếu đã có lab/data/filter_report.json và rr_malicious_kept = 0. Full ~15–40 phút. Payloads.csv mặc định, encoding latin-1.")
    code(doc, """cd /home/kali/Desktop/ATBMHTTT/lab
# smoke 200 dòng
../.venv/bin/python src/filter_mereani.py --limit 200 \\
  --output /tmp/mereani_smoke.csv --report /tmp/filter_smoke.json
# full — đúng số luận văn
../.venv/bin/python src/filter_mereani.py
# → data/mereani_exec_oracle.csv
# → data/filter_report.json""")
    kv(doc, "Mereani gốc 43.217 (15.149 độc / 28.068 lành). Giữ execute 10.824. Parser-only 3.542 (W1 = 23,4%). RR trên tập giữ = 0. Sai RR thì dừng.")

    h(doc, "4.3. B3 — P2 split kiểu Table 3 và 27 action", 2)
    code(doc, """cd /home/kali/Desktop/ATBMHTTT/lab
../.venv/bin/python src/split_table3.py
../.venv/bin/python src/action_alive.py
ls data/splits/detectors/train.csv data/splits/adversarial_agents/train.csv data/action_alive.csv""")
    kv(doc, "Seed 42. Detector undersample Benign = số Malicious. Agent split chỉ Malicious. File: data/splits/split_report.json.")

    h(doc, "4.4. B4 — Docker app + CRS PL1/PL2", 2)
    para(doc, "Image owasp/modsecurity-crs:nginx-alpine. App build từ lab/app/Dockerfile (node:20-alpine). Compose đã ghim 127.0.0.1.")
    code(doc, """cd /home/kali/Desktop/ATBMHTTT/lab
docker compose up -d --build
docker compose ps
curl -sS http://127.0.0.1:13000/health
curl -sS -o /dev/null -w 'app %{http_code}\\n' \\
  'http://127.0.0.1:13000/html?q=%3Cscript%3Ealert(1)%3C/script%3E'
curl -sS -o /dev/null -w 'pl1 %{http_code}\\n' \\
  'http://127.0.0.1:18080/html?q=%3Cscript%3Ealert(1)%3C/script%3E'
curl -sS -o /dev/null -w 'hello %{http_code}\\n' \\
  'http://127.0.0.1:18080/html?q=hello'
curl -sS -o /dev/null -w 'pl2 %{http_code}\\n' \\
  'http://127.0.0.1:18081/html?q=%3Cscript%3Ealert(1)%3C/script%3E'
../.venv/bin/python -m unittest tests.test_crs_stack -v""")
    kv(doc, "Health ok. App script HTTP 200. CRS PL1 script 403. PL1 hello 200. PL2 script 403. Unittest CRS xanh.")
    para(doc, "Sink app: /html?q= , /attr?q= , /js?q= và /purify/html , /purify/attr , /purify/js. Tắt: docker compose down.")

    h(doc, "4.5. B5 — P3 LSTM + PPO (chưa Oracle trong reward)", 2)
    para(doc, "Cần data/splits/. Lệch paper đã ghi: Adam (SGD kẹt 50% trên split JSDOM), pad 40. Default timesteps trong script là 25000; luận văn dùng 20000.")
    code(doc, """cd /home/kali/Desktop/ATBMHTTT/lab
../.venv/bin/python src/run_p3.py --phase all --timesteps 20000 \\
  --ppo-seeds 42,43,44 --max-eval 400""")
    para(doc, "Hoặc tách pha: --phase vocab rồi detector rồi ppo rồi eval.")
    kv(doc, "Checkpoint lab/runs/p3/lstm.pt, vocab.json, ppo_seed42.zip… LSTM acc ~0,991. PPO ER trung bình 3 seed × 20k = 0,75%. Panic nếu ER ~99% (nhầm encode nhánh A).")

    h(doc, "4.6. B6 — P4 Oracle-in-loop + Dueling DQN", 2)
    para(doc, "Cần runs/p3/lstm.pt. Reward R_exec: −2 hết độc JSDOM; −1 còn độc nhưng detector bắt; +10 chỉ khi lọt VÀ execute.")
    code(doc, """cd /home/kali/Desktop/ATBMHTTT/lab
../.venv/bin/python src/run_p4.py --phase all --algo ppo --seeds 42,43,44 \\
  --timesteps 20000 --max-eval 400
../.venv/bin/python src/train_ddqn.py --phase all --seeds 42,43,44 \\
  --timesteps 20000 --max-eval 400""")
    kv(doc, "PPO ER = TASR = 0,25% (sàn FN 1/400). Dueling DQN tb 1,08%. RR(E) = 0. Wilcoxon P3 vs P4 n=2 p=0,5 — chỉ mô tả. Bảng: lab/data/PPO_VS_DDQN.md.")
    cam(doc, "Viết DQN mạnh hơn Chen; pretent Wilcoxon có ý nghĩa thống kê.")

    h(doc, "4.7. B7 — P5 đo TASR (Docker đang lên)", 2)
    para(doc, "Nhãn: BLOCK (403) | BYPASS_NOEXEC (200, JS chết) | BYPASS_EXEC (200 và execute). TASR = số BYPASS_EXEC / N. Escape = chỉ HTTP 200.")
    code(doc, """cd /home/kali/Desktop/ATBMHTTT/lab
../.venv/bin/python src/measure_tasr.py --sink /html
../.venv/bin/python src/measure_tasr.py --sink /html --pairs --only pl1
../.venv/bin/python src/complete_p5.py --bases 40 --transfer-n 80
printf '%s\\n' '{"payload":"<script>alert(1)</script>"}' '{"payload":"<div>hello</div>"}' \\
  | node oracle/pw_check.mjs""")
    kv(doc, "CRS PL1/PL2 TASR = 0 trên catalog 84, chuỗi 2–3 action, 1120 Mereani decode, transfer PPO. App TASR ~0,845. DOMPurify TASR = 0, Δ = 1. Playwright: script true, div false.")
    cam(doc, "Gate BYPASS_EXEC trên PL1 không đạt → không huấn luyện RL trên CRS, không LLM generator.")

    h(doc, "5. NHÁNH A — artifact/ (data/10, LabelEncoder, ER ~99%)", 1)
    para(
        doc,
        "Chỉ khi muốn tái lập số paper. Không thay số nhánh B. CWD bắt buộc là artifact/Adversarial_RL_XSS khi gọi python src/… vì mutators.py đọc CSV theo thư mục hiện tại. "
        "Không dùng split JSDOM của lab. Không chạy prepare_dataset.py.",
    )

    h(doc, "5.1. A1 — Oracle FastAPI cổng 5555", 2)
    para(doc, "Giữ process này sống cả phiên nhánh A (kể cả job không Oracle vẫn gọi :5555 khi đo RR). Không --reload.")
    code(doc, """cd /home/kali/Desktop/ATBMHTTT
bash lab/src/run_pasini_faithful.sh oracle-up
# tương đương: cd artifact/Adversarial_RL_XSS
# ../../.venv/bin/python -m uvicorn app.main:app --host 127.0.0.1 --port 5555""")
    para(doc, "Smoke (terminal khác):")
    code(doc, """curl -sS -o /dev/null -w '%{http_code}\\n' http://127.0.0.1:5555/docs
cd /home/kali/Desktop/ATBMHTTT/artifact/Adversarial_RL_XSS
../../.venv/bin/python - <<'PY'
import sys
sys.path.insert(0, "src")
from utils.request_tools import do_xss_post_request
from utils.html_tools import is_same_dom
ep = "http://127.0.0.1:5555/vuln_backend/1.0/endpoint/"
abc = do_xss_post_request(ep, "abc")
print("div_malicious", not is_same_dom(do_xss_post_request(ep, "<div>hello</div>"), abc))
print("script_malicious", not is_same_dom(do_xss_post_request(ep, "<script>alert(1)</script>"), abc))
PY""")
    kv(doc, "docs HTTP 200. Cả div_malicious và script_malicious = True. <div> là độc theo Oracle DOM — đúng paper, khác JSDOM.")
    para(doc, "POST body json {\"payload\": [str]} (chuỗi trần = 422). Tắt: kill PID uvicorn, tránh pkill -f trùng argv wrapper.")

    h(doc, "5.2. A2 — Detector MLP / LSTM / CNN (Table 4)", 2)
    para(doc, "SGD lr=0.001, batch 16, 150 epoch, patience 10, embed 8, pad 30 (code artifact; paper viết pad 200).")
    code(doc, """cd /home/kali/Desktop/ATBMHTTT
bash lab/src/run_pasini_faithful.sh detectors""")
    kv(doc, "Test n=1802, seed 42: Precision 99,67% Recall 100% Accuracy/F1 99,83%. Cả ba mạng một bộ số (TP 901 FP 3 TN 898 FN 0). Checkpoint: artifact/.../runs/{lstm,mlp,cnn}/10/run_0/")

    h(doc, "5.3. A3 — PPO không Oracle, seed 42", 2)
    para(doc, "Paper: 250.000 bước, max 15 bước, 27 mutators, test ER n=901. Fast: 50.000 bước, n_eval_episodes=32. Số luận văn seed 42 là bản 250k.")
    para(doc, "LSTM đủ bảo vệ (wrapper):")
    code(doc, """cd /home/kali/Desktop/ATBMHTTT
# Nhanh ~5 phút
bash lab/src/run_pasini_faithful.sh ppo-lstm-fast
bash lab/src/run_pasini_faithful.sh eval-lstm \\
  artifact/Adversarial_RL_XSS/runs/lstm/10/run_0/adversarial_agent_fast/run_0/best_model.zip
# Đúng 250k
bash lab/src/run_pasini_faithful.sh ppo-lstm
bash lab/src/run_pasini_faithful.sh eval-lstm""")
    para(doc, "MLP và CNN cùng protocol (CWD artifact):")
    code(doc, """cd /home/kali/Desktop/ATBMHTTT/artifact/Adversarial_RL_XSS
PY=../../.venv/bin/python
for m in mlp cnn; do
  $PY src/train_adversarial_agent.py \\
    --trainset data/10/adversarial_agents/train.csv \\
    --valset data/10/adversarial_agents/val.csv \\
    --config_detector runs/$m/10/run_0/config.json \\
    --runs_folder adversarial_agent --seed 42 --timesteps 250000 \\
    --n_eval_episodes 32 --skip_env_check
  $PY src/test_adversarial_agent.py \\
    --testset data/10/adversarial_agents/test.csv \\
    --config_detector runs/$m/10/run_0/config.json \\
    --checkpoint runs/$m/10/run_0/adversarial_agent/run_0/best_model.zip --seed 42
done""")
    para(doc, "RR (cần :5555). RecursionError DOM thì giữ ER, không bịa RR:")
    code(doc, """cd /home/kali/Desktop/ATBMHTTT/artifact/Adversarial_RL_XSS
../../.venv/bin/python src/test_validity_mutated_dataset.py \\
  --dataset runs/lstm/10/run_0/adversarial_agent/run_0/empirical_study_set.csv \\
  --vocab data/10/vocabulary.csv --seed 42
../../.venv/bin/python src/analyze_validity.py \\
  --dataset runs/lstm/10/run_0/adversarial_agent/run_0/validity.csv --seed 42""")
    kv(doc, "Seed 42, 250k, test 901: LSTM ER 98,78% (paper 10 seed 98,62%); MLP 98,67%; CNN 98,89%. LSTM RR(E)=8,99% RR(V)=98,54% OR(V)=48,54% — TH2. CNN seed 42 ER vẫn ~99% nhưng RR(V) thấp (phương sai 1 seed).")

    h(doc, "5.4. A4 — RQ3 Oracle-in-loop (reward code −5)", 2)
    para(doc, "Paper PDF viết −2; code artifact −5. Oracle :5555 phải sống. MLP 50k dưới train (ER ~83,8%); cần 250k.")
    code(doc, """cd /home/kali/Desktop/ATBMHTTT
bash lab/src/run_pasini_faithful.sh ppo-lstm-oracle
# nhanh: ppo-lstm-oracle-fast

cd artifact/Adversarial_RL_XSS
PY=../../.venv/bin/python
for m in mlp cnn; do
  $PY src/train_adversarial_agent.py \\
    --trainset data/10/adversarial_agents/train.csv \\
    --valset data/10/adversarial_agents/val.csv \\
    --config_detector runs/$m/10/run_0/config.json \\
    --runs_folder adversarial_agent_oracle --seed 42 --timesteps 250000 \\
    --oracle_guided_reward --n_eval_episodes 32 --skip_env_check
done""")
    kv(doc, "Seed 42: LSTM 98,34% (paper 98,13%); MLP 97,89% (97,37%); CNN 96,00% (96,89%). RR ≈ 0. Câu được: RQ3 ER > 96% theo Oracle DOM. Câu cấm: coi đó là XSS còn chạy JavaScript.")

    h(doc, "5.5. A5 — Lưới 10 seed (optional)", 2)
    para(doc, "Paper công bố trung bình 10 seed × 3 mạng × 2 pha = 60 job, seeds 42–51, 250k bước. Seed 42 đủ Table 4 + RQ3. Máy này đã 29/60.")
    code(doc, """# :5555 sống cả lưới
cd /home/kali/Desktop/ATBMHTTT
.venv/bin/python -u lab/src/run_pasini_10seed.py
# chỉ xuất bảng, không train:
.venv/bin/python -u lab/src/run_offline_bao_cao.py --export-only""")
    para(doc, "Log: lab/runs/pasini_faithful/train_10seed.log. Tóm tắt: seed10_summary.json. Bảng dán: huong_dan/BANG_KET_QUA_THI_NGHIEM.txt.")

    h(doc, "6. Số kỳ vọng sau khi build đúng", 1)
    tbl(
        doc,
        ["Sau bước", "Kỳ vọng"],
        [
            ["unittest oracle", "<div> không execute; <script>alert(1) execute"],
            ["filter_report.json", "rr_malicious_kept = 0, W1 ≈ 23,4%"],
            ["curl PL1 <script>", "HTTP 403"],
            ["curl app <script>", "HTTP 200"],
            ["P3 PPO 3×20k", "ER tb 0,75% (panic nếu ~99%)"],
            ["P4 PPO Oracle", "ER = TASR = 0,25%; RR = 0"],
            ["P4 Dueling DQN", "1,08%; RR = 0"],
            ["P5 PL1 Table 2", "TASR = 0; Purify Δ = 1; app ~84,5%"],
            ["Nhánh A Table 4", "P 99,67 R 100 Acc 99,83 — 3 mạng trùng"],
            ["Nhánh A LSTM PPO", "ER 98,78%; RQ3 98,34% RR≈0"],
        ],
    )

    h(doc, "7. Lệnh chụp hình đưa vào báo cáo (không train)", 1)
    para(doc, "Terminal VS Code, font 16+. Không chạy filter_mereani full, run_p3/run_p4 train, lưới 10 seed. Docker CRS phải Up.")
    para(doc, "Một mạch có dừng Enter giữa shot:", bold=True)
    code(doc, """cd /home/kali/Desktop/ATBMHTTT
bash lab/src/demo_record.sh""")
    para(doc, "Từng shot (copy từng khối, chụp):")
    code(doc, """cd /home/kali/Desktop/ATBMHTTT/lab
# Shot 1 Oracle
node oracle/jsdom_oracle.mjs --payload '<script>alert(1)</script>'
node oracle/jsdom_oracle.mjs --payload '<div>hello</div>'
# Shot 2 unittest
../.venv/bin/python -m unittest tests.test_oracle tests.test_actions -v
# Shot 4 CRS
curl -sS -o /dev/null -w 'app  HTTP %{http_code}\\n' \\
  'http://127.0.0.1:13000/html?q=%3Cscript%3Ealert(1)%3C/script%3E'
curl -sS -o /dev/null -w 'PL1 HTTP %{http_code}\\n' \\
  'http://127.0.0.1:18080/html?q=%3Cscript%3Ealert(1)%3C/script%3E'
curl -sS -o /dev/null -w 'hello HTTP %{http_code}\\n' \\
  'http://127.0.0.1:18080/html?q=hello'""")
    para(doc, "Mở file (Ctrl+P), không Run: lab/data/PASINI_FAITHFUL_RESULTS.md, P3_RESULTS.md, PPO_VS_DDQN.md, P5_RESULTS.md, lab/runs/p3/p3_eval.json, artifact/.../xss_dataset.py (dòng fit_transform).")

    h(doc, "8. Thời gian CPU tham chiếu", 1)
    tbl(
        doc,
        ["Bước", "Thời gian"],
        [
            ["venv + pip + npm + playwright", "10–20 phút (mạng)"],
            ["unittest oracle/actions", "< 1 phút"],
            ["filter_mereani.py full", "15–40 phút"],
            ["docker compose up --build lần đầu", "5–15 phút"],
            ["P3 detector", "~2 phút"],
            ["P3 PPO 3 seed × 20k", "10–20 phút"],
            ["P4 PPO 3 seed + DDQN", "15–30 phút"],
            ["P5 measure_tasr 84 payload", "vài phút"],
            ["Nhánh A 3 detector", "~3 phút"],
            ["Nhánh A PPO 250k / mạng", "~13–15 phút không Oracle; ~25–32 phút Oracle"],
            ["Lưới 10 seed đủ 60 job", "khoảng nửa ngày–1 ngày"],
        ],
    )

    h(doc, "9. Lỗi thường gặp", 1)
    tbl(
        doc,
        ["Triệu chứng", "Sửa"],
        [
            ["jsdom missing", "cd lab && npm install"],
            ["latin-1 UnicodeDecodeError", "Payloads.csv phải encoding=latin-1 (script đã set)"],
            ["CRS SkipTest / connection refused", "docker compose up -d; đợi health; logs crs-pl1"],
            ["PL1 trả 200 cho script", "Đúng cổng 18080, không đo WAF trên 13000"],
            ["PPO ER ~99% trên lab JSDOM", "Nhầm nhánh A encode; nhánh B phải vocab.json cố định"],
            ["nn.Embedding(..., 8.0)", "artifact đã int(embedding_dim)"],
            ["mutators FileNotFoundError", "cd artifact/Adversarial_RL_XSS"],
            ["PPO quá lâu", "n_eval 32, skip_env_check; đừng eval 721 lúc train"],
            ["FastAPI 422", "POST {\"payload\": [str]} không phải chuỗi trần"],
            ["RR RecursionError", "Giữ ER; RR để trống"],
        ],
    )

    h(doc, "10. Checklist lab sống", 1)
    items = [
        ".venv import torch, pandas, gymnasium, stable_baselines3",
        "node oracle: script execute, div không",
        "unittest test_oracle + test_actions xanh",
        "filter_report.json rr_malicious_kept = 0 (hoặc file lọc đã có)",
        "data/splits/detectors/train.csv tồn tại",
        "docker compose ps: app, crs-pl1, crs-pl2 Up (healthy)",
        "curl app script 200; PL1 script 403; PL1 hello 200",
        "unittest test_crs_stack xanh",
        "(nhánh B) runs/p3/lstm.pt và p3_eval.json",
        "(nhánh B) runs/p4 và PPO_VS_DDQN.md",
        "(nhánh B) tasr_report.json TASR PL1 = 0",
        "(nhánh A) uvicorn :5555 + Table 4 test_results.json",
        "(nhánh A) results.json LSTM PPO + RQ3",
    ]
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        r = p.add_run()
        set_run(r, it, size=12)
        p.paragraph_format.space_after = Pt(2)

    h(doc, "11. Bản đồ file", 1)
    tbl(
        doc,
        ["Việc", "File"],
        [
            ["Sổ tay dựng máy", "huong_dan/HUONG_DAN_BUILD_LAB.md"],
            ["File Word này (generate)", "huong_dan/_gen_quytrinh.py → Quy_trinh_tai_hien_day_du.docx"],
            ["P1–P2 số", "lab/data/P1_P2_RESULTS.md"],
            ["P3 số", "lab/data/P3_RESULTS.md / runs/p3/p3_eval.json"],
            ["P4 / DDQN", "lab/data/P4_RESULTS.md / PPO_VS_DDQN.md"],
            ["P5 CRS", "lab/data/P5_RESULTS.md / tasr_report.json / p5_complete.json"],
            ["Nhánh A số", "lab/data/PASINI_FAITHFUL_RESULTS.md"],
            ["Wrapper nhánh A", "lab/src/run_pasini_faithful.sh"],
            ["10 seed", "lab/src/run_pasini_10seed.py"],
            ["Demo chụp/quay", "lab/src/demo_record.sh"],
            ["Báo cáo dán", "huong_dan/BAO_CAO_DO_AN_FULL.txt"],
            ["Slide", "huong_dan/Slide_Bao_ve_ATBMHTTT.pptx"],
        ],
    )

    h(doc, "12. Câu được / câu cấm khi viết báo cáo", 1)
    para(doc, "Được:", bold=True, space_after=4)
    for t in [
        "Khung Pasini: Oracle → Ruin Rate → Oracle-in-the-loop.",
        "23,4% nhãn Malicious Mereani chỉ đổi DOM.",
        "Trên đúng artifact data/10, tái lập ER ~99% và RQ3 ER > 96% với RR ≈ 0 theo Oracle DOM.",
        "Khi khép token-id, PPO Table 2 không còn ER ~99%.",
        "CRS lab catalog Table 2 TASR = 0. DOMPurify Δ = 1.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        r = p.add_run()
        set_run(r, t, size=12)
    para(doc, "Cấm:", bold=True, color=RED, space_before=8, space_after=4)
    for t in [
        "Em reproduce Chen 99% là thành công của đồ án.",
        "Replication SAC; PPO/DQN mạnh hơn Chen.",
        "Bypass CRS / AWS WAF.",
        "HTTP 200 = TASR / thành công.",
        "LLM generator; Wilcoxon p=0,5 là có ý nghĩa.",
        "Trộn bảng nhánh A với nhánh B.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        r = p.add_run()
        set_run(r, t, size=12)

    para(
        doc,
        "Hết quy trình. Repo gốc /home/kali/Desktop/ATBMHTTT. Sinh lại file này: python3 huong_dan/_gen_quytrinh.py",
        italic=True,
        size=12,
        space_before=16,
    )

    doc.save(str(OUT))
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
